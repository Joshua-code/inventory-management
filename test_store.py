import json
import os
import struct
import tempfile

import openpyxl
from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
from cryptography.hazmat.primitives.serialization import Encoding, NoEncryption, PrivateFormat, PublicFormat

import keygen
import licensing
import updater
from app import FILE_MAX, LABEL_W, db_info_text, label_bytes, parse_margin, render_label, saved_margins
from store import PERMS, AccessRevoked, Store, _perms, parse_excel

d = tempfile.mkdtemp()
xlsx = os.path.join(d, "barang.xlsx")
wb = openpyxl.Workbook()
wb.active.append(["Nama Barang", "Harga", "Barcode"])
wb.active.append(["Indomie Goreng", 3500, 8998866200301])
wb.active.append(["Aqua 600ml", "Rp 4.000", "ABC-1"])
wb.active.append([None, None, None])
wb.save(xlsx)

db_dir = os.path.join(d, "data")
os.mkdir(db_dir)
path = os.path.join(db_dir, "toko.rptdb")
s = Store(path)
assert not s.exists()
s.setup("admin", "rahasia")
assert s.perms == [] and s.login("admin", "salah") is None and s.login("nobody", "rahasia") is None
assert s.login("admin", "rahasia") == list(PERMS)
assert s.import_excel(xlsx) == 2 and s.source == "barang.xlsx"
assert s.imported and s.imported == Store(path).info()[1]  # encrypted copy matches the header
assert s.lookup("8998866200301") == ["Indomie Goreng", 3500]
assert s.lookup(" ABC-1 ") == ["Aqua 600ml", 4000]
s.add_user("kasir", "123", ["scanner"])
s.add_user("tmp", "456", ["printer", "scanner"])
assert s.list_users() == [("admin", list(PERMS)), ("kasir", ["scanner"]), ("tmp", ["scanner", "printer"])]
for bad_args in [("x", "x", []), ("x", "x", ["nope"]), ("admin", "x", ["scanner"])]:
    try:
        s.add_user(*bad_args)
        raise AssertionError(bad_args)
    except ValueError:
        pass
s.set_perms("tmp", ["printer"])
assert dict(s.list_users())["tmp"] == ["printer"]
for bad_args in [("tmp", []), ("admin", ["scanner", "update"]), ("ghost", ["scanner"])]:
    try:
        s.set_perms(*bad_args)
        raise AssertionError(bad_args)
    except ValueError:
        pass
assert _perms("admin") == list(PERMS) and _perms("scanner") == ["scanner"] and _perms("printer") == ["printer"]
assert _perms(["users", "bogus", "scanner"]) == ["scanner", "users"] and _perms(None) == []

bad = os.path.join(d, "bad.xlsx")
wb.active.append(["Teh", "mahal", "999"])
wb.save(bad)
try:
    s.import_excel(bad)
    raise AssertionError("bad excel accepted")
except ValueError as e:
    assert "Baris 5" in str(e)
assert s.lookup("ABC-1") and s.source == "barang.xlsx"  # old data kept after failed import

s.delete_user("tmp")
assert Store(path).login("tmp", "456") is None

# exactly one file, nothing readable inside it
assert os.listdir(db_dir) == ["toko.rptdb"]
raw = open(path, "rb").read()
for secret in [b"admin", b"kasir", b"scanner", b"Indomie", b"8998866200301"]:
    assert secret not in raw, secret

file, when = Store(path).info()  # readable before login
assert file == "barang.xlsx" and when and Store(os.path.join(d, "none.rptdb")).info() == (None, None)

s.logout()
s2 = Store(path)
assert s2.login("kasir", "123") == ["scanner"]
assert s2.lookup("8998866200301") and s2.source == "barang.xlsx"
for fn, args in [(s2.add_user, ("x", "x", ["scanner"])), (s2.import_excel, (xlsx,)), (s2.list_users, ())]:
    try:
        fn(*args)
        raise AssertionError(f"{fn.__name__} without access")
    except PermissionError:
        pass

# sync: another laptop changes the file -> reload picks it up without logging out
assert s2.reload_if_changed() is False
admin = Store(path)
admin.login("admin", "rahasia")
wb2 = openpyxl.Workbook()
wb2.active.append(["Kopi Kapal Api", 1500, "555"])
x2 = os.path.join(d, "baru.xlsx")
wb2.save(x2)
admin.import_excel(x2)
os.utime(path, (100, 100))  # distinct mtime per write, even on coarse (Windows) clocks
assert s2.reload_if_changed() is True and s2.lookup("555") == ["Kopi Kapal Api", 1500] and s2.source == "baru.xlsx"
assert s2.imported == admin.imported
assert s2.reload_if_changed() is False
file, when = Store(path).info()  # footer info follows the latest import
assert file == "baru.xlsx" and when
admin.set_perms("kasir", ["scanner", "printer"])
os.utime(path, (101, 101))
assert s2.reload_if_changed() and s2.perms == ["scanner", "printer"]
admin.delete_user("kasir")
os.utime(path, (102, 102))
try:
    s2.reload_if_changed()
    raise AssertionError("deleted user kept access")
except AccessRevoked:
    assert s2.key is None
raw = open(path, "rb").read()

# tampering with the encrypted data must not open
doc = json.loads(raw)
doc["data"] = doc["data"][:-2] + ("00" if doc["data"][-2:] != "00" else "11")
json.dump(doc, open(path, "w"))
try:
    Store(path).login("admin", "rahasia")
    raise AssertionError("tampered data opened")
except Exception as e:
    assert not isinstance(e, AssertionError)

# not a database
junk = os.path.join(d, "junk.rptdb")
open(junk, "w").write("hello")
try:
    Store(junk).login("admin", "x")
    raise AssertionError("junk accepted")
except ValueError as e:
    assert "valid" in str(e)

# header is optional: no header -> row 1 is data; any header text (Harga not a number) -> skipped
for first, n in [(["Teh Botol", 5000, "111"], 2), (["Nama", "Harga Jual", "Kode"], 1)]:
    wb = openpyxl.Workbook()
    wb.active.append(first)
    wb.active.append(["Kopi", 2000, "222"])
    p = os.path.join(d, "h.xlsx")
    wb.save(p)
    assert len(parse_excel(p)) == n, first

# import date: encrypted copy wins, the plaintext header only covers older databases
t = Store(path)
t.user = "admin"
t._load({"items": {}, "file": None, "users": {"admin": ["scanner"]}}, {"imported": "01/01/2026 10:00"})
assert t.imported == "01/01/2026 10:00"
t._load({"items": {}, "file": None, "imported": "02/02/2026 11:00", "users": {"admin": ["scanner"]}}, {"imported": "01/01/2026 10:00"})
assert t.imported == "02/02/2026 11:00"


def raster_rows(data):
    assert data.startswith(b"\x1b@") and b"\x1bJ" not in data and b"\x1dV" not in data  # no feed/cut commands
    i, rows = 2, 0
    while i < len(data):
        assert data[i:i + 4] == b"\x1dv0\x00"
        xb, yb = struct.unpack("<HH", data[i + 4:i + 8])
        assert xb == LABEL_W // 8
        rows += yb
        i += 8 + xb * yb
    return rows


# label: margins are blank rows, the import date shows only when given
img = render_label("Indomie Goreng Rasa Ayam Bawang Special", 125000, "8998866200301")
dated = render_label("Indomie Goreng Rasa Ayam Bawang Special", 125000, "8998866200301", "29/09/2026 09:50")
assert img.width == LABEL_W and dated.size == img.size  # short barcode: date shares the number's line


def ink(im, box):
    return sum(im.crop(box).histogram()[:128])  # pixels darker than mid-grey


corner = (LABEL_W - 100, img.height - 24, LABEL_W, img.height)  # bottom-right, beside the barcode number
assert ink(img, corner) == 0 and ink(dated, corner) > 0
long_code = render_label("Teh", 5000, "ABCDEFGHIJKLMNOPQRST", "29/09/2026 09:50")  # number too wide: date goes below
assert long_code.height > render_label("Teh", 5000, "ABCDEFGHIJKLMNOPQRST").height
for top, bottom in [(0, 0), (0, 25), (3, 60), (60, 60)]:
    assert raster_rows(label_bytes(img, top, bottom)) == img.height + (top + bottom) * 8, (top, bottom)
for text, ok in [("0", 0), (" 25 ", 25), ("60", 60)]:
    assert parse_margin(text) == ok
for bad_text in ["-1", "61", "abc", "", "2.5"]:
    try:
        parse_margin(bad_text)
        raise AssertionError(bad_text)
    except ValueError:
        pass
assert saved_margins({}) == (0, 25) and saved_margins({"margin_top_mm": 3, "margin_bottom_mm": "40"}) == (3, 40)
assert saved_margins({"margin_top_mm": "x", "margin_bottom_mm": 999}) == (0, 25)

# footer: database file + import date; long names shortened, extension kept
assert db_info_text(None, None) == ""
assert db_info_text("barang.xlsx", "08/10/2026 15:18") == "File: barang.xlsx  ·  Diimpor 08/10/2026 15:18"
long_name = "daftar_harga_barang_semua_cabang_babel_mart_oktober_2026_final.xlsx"
short = db_info_text(long_name, "x").removeprefix("File: ").removesuffix("  ·  Diimpor x")
assert len(short) == FILE_MAX and short.startswith("daftar_harga") and short.endswith("2026_final.xlsx") and "…" in short

# license: code signed for device A works only on A
k = Ed25519PrivateKey.generate()
licensing.PUBLIC_KEY = k.public_key().public_bytes(Encoding.Raw, PublicFormat.Raw).hex()
code = licensing.format_code(k.sign(b"AAAA-BBBB-CCCC-DDDD"))
assert licensing.verify(code, "AAAA-BBBB-CCCC-DDDD")
assert licensing.verify(" " + code.lower().replace("-", " ") + "\n", "AAAA-BBBB-CCCC-DDDD")
assert not licensing.verify(code, "AAAA-BBBB-CCCC-DDDE")
assert not licensing.verify("garbage", "AAAA-BBBB-CCCC-DDDD") and not licensing.verify("", "AAAA-BBBB-CCCC-DDDD")
assert len(licensing.device_id()) == 19

# generator: only the matching key file is accepted; ID format is checked
raw = lambda key: key.private_bytes(Encoding.Raw, PrivateFormat.Raw, NoEncryption())  # noqa: E731
for name, data in [("good.key", raw(k)), ("other.key", raw(Ed25519PrivateKey.generate())), ("junk.key", b"x")]:
    open(os.path.join(d, name), "wb").write(data)
gk = keygen.load_key(os.path.join(d, "good.key"))
for name in ["other.key", "junk.key"]:
    try:
        keygen.load_key(os.path.join(d, name))
        raise AssertionError(name)
    except ValueError:
        pass
assert licensing.verify(keygen.make_code(gk, "aaaabbbbccccdddd"), "AAAA-BBBB-CCCC-DDDD")
try:
    keygen.make_code(gk, "not-an-id")
    raise AssertionError("bad id accepted")
except ValueError:
    pass

# updates: only a newer, signed Setup whose hash matches is accepted
folder = os.path.join(d, "drive")
os.mkdir(folder)
setup = os.path.join(folder, "Retail Price Tag Management Setup 1.2.0.exe")
open(setup, "wb").write(b"MZ-pretend-installer")
manifest = updater.write_manifest(k, setup)
info = updater.check(folder, "1.1.0")
assert info and info["version"] == "1.2.0" and info["file"] == os.path.basename(setup)
assert updater.check(folder, "1.2.0") is None and updater.check(folder, "1.10.0") is None  # not newer
assert updater.parse_version("1.0.10") > updater.parse_version("1.0.9") and updater.check(os.path.join(d, "none"), "1.0.0") is None
good = open(manifest).read()
m = json.loads(good)
for field, value in [("version", "1.3.0"), ("sha256", "0" * 64), ("sig", "AAAA"), ("sha256", "xyz"), ("file", "../evil.exe")]:
    json.dump({**m, field: value}, open(manifest, "w"))
    assert updater.check(folder, "1.1.0") is None, field
open(manifest, "w").write("not json")
assert updater.check(folder, "1.1.0") is None
open(manifest, "w").write(good)
assert updater.check(folder, "1.1.0")
tmp_setup = updater.prepare(info, folder)  # verified private copy
assert tmp_setup != setup and updater.sha256_file(tmp_setup) == info["sha256"]
cmd = updater.install_cmd(tmp_setup)
assert cmd[0] == tmp_setup and "/SILENT" in cmd and "/RELAUNCH=1" in cmd
open(setup, "ab").write(b"tampered")  # swapped after the manifest was signed
try:
    updater.prepare(info, folder)
    raise AssertionError("modified Setup accepted")
except ValueError:
    pass
other_key = Ed25519PrivateKey.generate()  # a manifest signed by anyone else is ignored
updater.write_manifest(other_key, setup)
assert updater.check(folder, "1.1.0") is None
try:
    updater.write_manifest(k, os.path.join(d, "setup-without-version.exe"))
    raise AssertionError("no version in file name")
except (ValueError, OSError):
    pass
print("OK")
